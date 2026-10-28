"""Step 4: write the colour-coded reconciliation workbook (openpyxl)."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .results import STATUS_HELP, STATUS_ORDER
from .schema import TermSheetExtraction

log = logging.getLogger(__name__)

NAVY = "1F3864"
STATUS_FILL = {
    "Match": "C6EFCE",
    "False Alert": "FFEB9C",
    "True Mismatch": "FFC7CE",
    "Parsing Error": "E4DFEC",
    "Not in Term Sheet": "EDEDED",
    "Missing in Booking": "FCE4D6",
    "Not Compared": "F2F2F2",
    "No Term Sheet": "F8CBAD",
    "Outlier (Review)": "FFD966",
    "Correct": "C6EFCE",
    "Wrong value": "FFC7CE",
    "Missed (returned null)": "FCE4D6",
    "Hallucinated (should be null)": "FFC7CE",
    "Parsing error": "E4DFEC",
    "BREAK": "FFC7CE",
    "REVIEW": "FFEB9C",
    "CLEAN": "C6EFCE",
    "UNLINKED": "F8CBAD",
    "OK": "C6EFCE",
    "NOT_FOUND": "EDEDED",
    "PARSING_ERROR": "E4DFEC",
    "ERROR": "FFC7CE",
    "WARNING": "FFEB9C",
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
WRAP = Alignment(wrap_text=True, vertical="top")


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _write_table(
    ws: Worksheet,
    df: pd.DataFrame,
    start_row: int = 1,
    status_cols: tuple[str, ...] = (),
    widths: dict[str, int] | None = None,
    freeze: bool = True,
) -> int:
    widths = widths or {}
    cols = list(df.columns)
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=start_row, column=j, value=c)
        cell.font, cell.fill, cell.border = HEADER_FONT, HEADER_FILL, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for i, rec in enumerate(df.itertuples(index=False), start_row + 1):
        for j, (c, v) in enumerate(zip(cols, rec, strict=False), 1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            elif v is not None and not isinstance(v, (str, int, float, bool)):
                v = str(v)
            cell = ws.cell(row=i, column=j, value=v)
            cell.border, cell.alignment = BORDER, WRAP
            if c in status_cols and isinstance(v, str) and v in STATUS_FILL:
                cell.fill = _fill(STATUS_FILL[v])
                cell.font = Font(bold=True)
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = widths.get(c, min(max(len(str(c)) + 2, 12), 40))
    last = start_row + len(df)
    if len(df):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{last}"
    if freeze:
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return last


def _summary_sheet(ws: Worksheet, detail: pd.DataFrame, trades: pd.DataFrame, meta: dict) -> None:
    ws.title = "Summary"
    ws["A1"] = "Term Sheet vs Booking Reconciliation"
    ws["A1"].font = Font(bold=True, size=16, color=NAVY)
    r = 3
    for k, v in meta.items():
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=str(v))
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Trade-level outcome").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    trade_counts = trades["overall_status"].value_counts() if not trades.empty else pd.Series(dtype=int)
    tdf = pd.DataFrame(
        [
            {"Trade status": s, "Trades": int(trade_counts.get(s, 0))}
            for s in ["BREAK", "REVIEW", "CLEAN", "UNLINKED"]
            if trade_counts.get(s, 0) or s != "UNLINKED"
        ]
    )
    r = _write_table(ws, tdf, r, status_cols=("Trade status",), freeze=False) + 2

    ws.cell(row=r, column=1, value="Field-level outcome by term sheet").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    pivot = detail.assign(document=detail["document"].fillna("-")).pivot_table(
        index="final_status", columns="document", values="booking_ref", aggfunc="count", fill_value=0
    )
    pivot["Total"] = pivot.sum(axis=1)
    pivot = (
        pivot.reindex([s for s in STATUS_ORDER if s in pivot.index])
        .reset_index()
        .rename(columns={"final_status": "Status"})
    )
    pivot.insert(1, "Meaning", pivot["Status"].map(STATUS_HELP))
    r = _write_table(ws, pivot, r, status_cols=("Status",), widths={"Meaning": 70, "Status": 20}, freeze=False) + 2

    ws.cell(row=r, column=1, value="How to read this workbook").font = Font(bold=True, size=12, color=NAVY)
    notes = [
        "Trade Summary - one row per booked trade; BREAK = at least one True Mismatch.",
        "Exceptions - only rows needing action (True Mismatch, Parsing Error, Missing in Booking, items flagged for review).",
        "Reconciliation Detail - every trade x field with term sheet value, booking value, verdict, reason and the verbatim term sheet evidence.",
        "Term Sheet Extraction - raw LLM extraction per document, schema status, evidence quote and whether the quote was found verbatim in the source.",
        "Extraction Accuracy - each extracted field scored against hand-verified golden values (data/expected).",
        "Booking Data - booking extracts after PySpark normalisation and CSV/JSON de-duplication.",
        "Data Quality - ingestion issues and extraction warnings.",
    ]
    for n in notes:
        r += 1
        ws.cell(row=r, column=1, value="• " + n)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 70


def write_workbook(
    path: Path,
    *,
    detail: pd.DataFrame,
    trades: pd.DataFrame,
    bookings: pd.DataFrame,
    extractions: list[TermSheetExtraction],
    dq_issues: list[dict],
    meta: dict,
    accuracy: pd.DataFrame | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _summary_sheet(wb.active, detail, trades, meta)

    ws = wb.create_sheet("Trade Summary")
    _write_table(
        ws,
        trades,
        status_cols=("overall_status",),
        widths={"booking_ref": 22, "source_files": 30, "true_mismatch_fields": 45, "parsing_error_fields": 30},
    )

    det_cols = {
        "document": "Term Sheet",
        "booking_ref": "Booking Ref",
        "TradeID": "Trade ID",
        "source_files": "Booking Source",
        "field_label": "Field",
        "booking_column": "Booking Column",
        "ts_value": "Term Sheet Value",
        "booking_value": "Booking Value",
        "initial_result": "Initial Comparison",
        "final_status": "Final Status",
        "reason": "Reason",
        "ts_quote": "Term Sheet Evidence (verbatim)",
        "quote_verified": "Evidence Verified",
        "ts_note": "Extraction Note",
        "llm_confidence": "LLM Confidence",
        "adjudicated_by": "Adjudicated By",
        "needs_review": "Needs Review",
        "extraction_method": "Extraction Method",
        "link_method": "Linked By",
    }
    detail_view = detail[[c for c in det_cols if c in detail.columns]].rename(columns=det_cols)
    widths = {
        "Reason": 60,
        "Term Sheet Evidence (verbatim)": 60,
        "Extraction Note": 40,
        "Field": 26,
        "Term Sheet Value": 22,
        "Booking Value": 22,
        "Booking Source": 28,
        "Booking Ref": 20,
        "Term Sheet": 24,
    }

    exceptions = detail_view[
        detail_view["Final Status"].isin(["True Mismatch", "Parsing Error", "Missing in Booking", "No Term Sheet"])
        | detail_view["Needs Review"]
    ]
    ws = wb.create_sheet("Exceptions")
    _write_table(ws, exceptions, status_cols=("Final Status",), widths=widths)

    ws = wb.create_sheet("Reconciliation Detail")
    _write_table(ws, detail_view, status_cols=("Final Status",), widths=widths)

    ex_rows = []
    for ex in extractions:
        for f in ex.fields.values():
            ex_rows.append(
                {
                    "Term Sheet": ex.document,
                    "Field": f.label,
                    "Schema Field": f.name,
                    "Status": f.status.value,
                    "Value": f.canonical(),
                    "Raw LLM Value": None if f.raw_value is None else str(f.raw_value),
                    "Evidence (verbatim quote)": f.quote,
                    "Quote Found in Source": f.quote_verified,
                    "Confidence": f.confidence,
                    "Note": f.note,
                    "Validation Error": f.error,
                    "Extraction Method": ex.extraction_method,
                }
            )
    ws = wb.create_sheet("Term Sheet Extraction")
    _write_table(
        ws,
        pd.DataFrame(ex_rows),
        status_cols=("Status",),
        widths={"Evidence (verbatim quote)": 70, "Note": 45, "Validation Error": 40, "Value": 24, "Raw LLM Value": 24},
    )

    if accuracy is not None and not accuracy.empty:
        ws = wb.create_sheet("Extraction Accuracy")
        acc = accuracy.drop(columns=["correct"]).rename(
            columns={
                "document": "Term Sheet",
                "field": "Field",
                "schema_field": "Schema Field",
                "expected": "Expected (golden)",
                "extracted": "Extracted",
                "outcome": "Outcome",
                "quote_verified": "Quote Found in Source",
                "extraction_method": "Extraction Method",
            }
        )
        _write_table(
            ws, acc, status_cols=("Outcome",), widths={"Expected (golden)": 26, "Extracted": 26, "Outcome": 26}
        )

    ws = wb.create_sheet("Booking Data")
    _write_table(ws, bookings, widths={"source_files": 30, "booking_ref": 22})

    dq = pd.DataFrame(dq_issues, columns=["source_file", "row", "column", "severity", "issue"])
    for ex in extractions:
        for w in ex.warnings:
            dq.loc[len(dq)] = [ex.document, None, None, "WARNING", w]
    if dq.empty:
        dq.loc[0] = ["-", None, None, "OK", "No data-quality issues detected"]
    ws = wb.create_sheet("Data Quality")
    _write_table(ws, dq, status_cols=("severity",), widths={"issue": 80, "source_file": 30, "row": 30})

    wb.save(path)
    log.info("Wrote reconciliation workbook: %s", path)
    return path


def default_report_name() -> str:
    return f"reconciliation_report_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
