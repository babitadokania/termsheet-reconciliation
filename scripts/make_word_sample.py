"""Create samples/word/Term_Sheet_Genel_Energy_KeyTerms.docx.

A Word rendition of the key terms on page 1 of the supplied Genel Energy PDF term sheet, laid out as
a two-column table (the typical Word term sheet layout). Used to demonstrate and test .docx support.

    python scripts/make_word_sample.py
"""

from pathlib import Path

import docx
from docx.shared import Pt

ROWS = [
    ("ISIN", "NO 0010894330"),
    ("Issuer", "Genel Energy Finance 4 Limited (incorporated in England, company registration number 11867181)"),
    ("Parent", "Genel Energy plc (incorporated in Jersey, company registration number 107897)."),
    ("Currency", "USD"),
    ("Issue Amount", "USD 300,000,000"),
    ("Coupon Rate", "9.25% p.a., semi-annual interest payments."),
    ("Issue Price", "97% of the Nominal Amount."),
    ("Issue Date", "Expected to be 14 October 2020."),
    ("Maturity Date", "Expected to be 14 October 2025 (5 years after Issue Date)."),
    ("First Interest Payment Date", "Expected to be 14 April 2021 (6 months after Issue Date)."),
    (
        "Interest Payments",
        "Interest on the Bonds will start to accrue on Issue Date and shall be payable semi-annually "
        "in arrears. Day-count fraction for coupon is “30/360”, business day convention is "
        "“unadjusted” and business day is “Oslo” and “New York”.",
    ),
    ("Amortization", "The Bonds shall be repaid in full at the Maturity Date at 100% of the Nominal Amount."),
    ("Nominal Amount", "The Bonds will have a nominal amount of USD 2,000 each."),
    ("Minimum Subscription", "Minimum subscription and allotment amount shall be USD 200,000."),
]


def main() -> Path:
    out = Path(__file__).resolve().parents[1] / "samples" / "word" / "Term_Sheet_Genel_Energy_KeyTerms.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    d.add_heading("Term Sheet - Genel Energy Finance 4 Limited", level=1)
    p = d.add_paragraph("Senior Unsecured Guaranteed Callable Bond Issue 2020/2025 (key terms)")
    p.runs[0].font.size = Pt(10)
    table = d.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for k, v in ROWS:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, v
    d.save(out)
    print(f"Wrote {out}")
    return out


if __name__ == "__main__":
    main()
